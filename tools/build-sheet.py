"""Build a map-ready copy of the Waterbrooke missionary workbook.

Keeps the existing 'Missionary Data' and 'Instructions & Notes' tabs, adds the
map columns J..O to Missionary Data (pre-filled where the Location text allows),
adds the lookup tabs from the generated seeds, adds a derived 'Data' tab (the
only tab to publish), a Country dropdown, and red-cell rules.

Needs: Python 3 and openpyxl (pip install openpyxl).

Usage:
    python build-sheet.py <downloaded-sheet.xlsx> [output-dir]

Download the shared sheet first, for example:
    curl -L -o sheet.xlsx "https://docs.google.com/spreadsheets/d/<ID>/export?format=xlsx"

The output contains private data (respondent names, prayer requests).
Keep it OUT of this repository. The default output directory is
../org-heatmap-sheet next to the repository.
"""
import csv
import os
import sys
import unicodedata

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

if len(sys.argv) < 2:
    sys.exit(__doc__)

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
SRC = sys.argv[1]
OUT_DIR = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(REPO), "org-heatmap-sheet")
OUT = os.path.join(OUT_DIR, "Waterbrooke Missionary Data (map-ready).xlsx")
SEEDS = os.path.join(REPO, "sheet")
MAX_ROW = 300  # formulas and rules cover rows 2..MAX_ROW

# Pre-fill plan keyed by the Organization text exactly as it appears in the sheet.
# label=None keeps the Organization text as the public Map Label.
PLAN = {
    "CRU / Danielle Demorret": dict(label="CRU", cc="US", state="Minnesota"),
    "Compassion International": dict(label=None, cc=None, state=None),
    "Calvary Chapel / Good Sports Hungary": dict(label=None, cc="HU", state=None),
    "RTS Heidelberg / Pastor Victor d'Assenville": dict(label="RTS Heidelberg", cc="DE", state="Baden-Württemberg"),
    "Brennan & Becca McCafferty / Mission to the World": dict(label="Mission to the World", cc="KH", state=None),
    "Bill & Kathy Craver / Direct Christian Impact": dict(label="Direct Christian Impact", cc="MX", state="Yucatán"),
    "Ross & Rossy Swihart / Mission Edu-cate": dict(label="Mission Edu-cate", cc="MX", state="Quintana Roo"),
    "Fellowship of Christian Athletes (FCA)": dict(label=None, cc="US", state="Wisconsin"),
    "InterCP": dict(label=None, cc="KR", state=None),
    "Voice of the Martyrs": dict(label=None, cc="US", state=None),
    "Envision Twin Cities": dict(label=None, cc="US", state="Minnesota"),
    "World Gospel Outreach / Rancho Ebenezer": dict(label=None, cc="HN", state="Francisco Morazán"),
    "Janet Myers / Kenya Real": dict(label="Kenya Real", cc="KE", state=None),
    "Gary & Lyndell Moberg": dict(label=None, cc="TH", state=None),
    "Tom Steller / Training Leaders International": dict(label="Training Leaders International", cc="CM", state=None),
}


def fold(s):
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower().strip()


def read_seed(name):
    with open(os.path.join(SEEDS, name), encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    return rows[0][0], rows[1], rows[2:]


def safe(v):
    """Keep strings that look like formulas as plain text."""
    return (" " + v) if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") else v


stamp_c, hdr_c, countries = read_seed("Countries.csv")
stamp_s, hdr_s, subs = read_seed("Subdivisions.csv")
canon_name = {r[2]: r[1] for r in countries if r[3] == "TRUE"}
sub_keys = {r[0] for r in subs}

wb = openpyxl.load_workbook(SRC)
ws = wb["Missionary Data"]

# ---------------------------------------------------------------------------
# Missionary Data: new columns J..O
# ---------------------------------------------------------------------------
NEW_HEADERS = ["Map Label", "Country", "State / Province", "Country Code", "Subdivision Code", "Region Code"]
FIRST_NEW = 10  # column J
hdr_src = ws.cell(row=1, column=9)
for i, h in enumerate(NEW_HEADERS):
    c = ws.cell(row=1, column=FIRST_NEW + i, value=h)
    c.font = hdr_src.font.copy()
    c.fill = hdr_src.fill.copy()
    c.alignment = hdr_src.alignment.copy()
    c.border = hdr_src.border.copy()
for col, width in zip("JKLMNO", (30, 22, 20, 12, 15, 12)):
    ws.column_dimensions[col].width = width

grey = Font(color="FF666666", italic=True)
report = []
for r in range(2, MAX_ROW + 1):
    src = ws.cell(row=r, column=9)
    for i in range(len(NEW_HEADERS)):
        cell = ws.cell(row=r, column=FIRST_NEW + i)
        if src.fill is not None and src.fill.fill_type:
            cell.fill = src.fill.copy()
    org = ws.cell(row=r, column=2).value
    plan = PLAN.get((org or "").strip()) if org else None
    if plan:
        label = plan["label"] or org.strip()
        ws.cell(row=r, column=10, value=label)
        country_name = canon_name.get(plan["cc"]) if plan["cc"] else None
        if country_name:
            ws.cell(row=r, column=11, value=country_name)
        state_note = ""
        if plan["state"]:
            key = f"{plan['cc'].lower()}|{plan['state'].lower()}"
            if key not in sub_keys:
                alt = f"{plan['cc'].lower()}|{fold(plan['state'])}"
                key = alt if alt in sub_keys else None
            if key:
                ws.cell(row=r, column=12, value=plan["state"])
            else:
                state_note = f" (state '{plan['state']}' NOT in lookup; left blank)"
        report.append((r, org.strip(), label, country_name or "(none: not placed)", plan["state"] or "", state_note))
    ws.cell(row=r, column=13, value=f'=IF(LEN(TRIM(K{r}))=0,"",IFERROR(VLOOKUP(LOWER(TRIM(K{r})),Countries!$A:$C,3,FALSE),""))').font = grey
    ws.cell(row=r, column=14, value=f'=IF(OR(LEN(TRIM(L{r}))=0,LEN(M{r})=0),"",IFERROR(VLOOKUP(LOWER(M{r})&"|"&LOWER(TRIM(L{r})),Subdivisions!$A:$D,4,FALSE),""))').font = grey
    ws.cell(row=r, column=15, value=f'=IF(LEN(M{r})=0,"",IF(LEN(N{r})>0,N{r},M{r}))').font = grey

# ---------------------------------------------------------------------------
# Lookup tabs
# ---------------------------------------------------------------------------
wc = wb.create_sheet("Countries")
wc.append([stamp_c])
wc.append(hdr_c + ["", "Dropdown"])
canon_sorted = sorted((r[1] for r in countries if r[3] == "TRUE"), key=str.lower)
for r in countries:
    wc.append([safe(x) for x in r])
for i, name in enumerate(canon_sorted, start=3):
    wc.cell(row=i, column=6, value=name)
wc.freeze_panes = "A3"
for col, width in zip("ABCDEF", (28, 30, 9, 10, 3, 30)):
    wc.column_dimensions[col].width = width
wc["A1"].font = grey
for c in wc[2]:
    c.font = Font(bold=True)

wsub = wb.create_sheet("Subdivisions")
wsub.append([stamp_s])
wsub.append(hdr_s)
for r in subs:
    wsub.append([safe(x) for x in r])
wsub.freeze_panes = "A3"
for col, width in zip("ABCDE", (34, 12, 34, 12, 10)):
    wsub.column_dimensions[col].width = width
wsub["A1"].font = grey
for c in wsub[2]:
    c.font = Font(bold=True)

# Country dropdown on Missionary Data!K
dropdown_last = 2 + len(canon_sorted)
dv = DataValidation(
    type="list", formula1=f"=Countries!$F$3:$F${dropdown_last}", allow_blank=True,
    showErrorMessage=True, errorTitle="Unknown country",
    error="Pick a country from the list. The names are the ones the map knows.",
)
ws.add_data_validation(dv)
dv.add(f"K2:K{MAX_ROW}")

# Red-cell rules on Missionary Data!K and L
red = PatternFill(start_color="FFF4CCCC", end_color="FFF4CCCC", fill_type="solid")
ws.conditional_formatting.add(f"K2:K{MAX_ROW}", FormulaRule(formula=["AND(LEN(TRIM(K2))>0,LEN(M2)=0)"], fill=red, stopIfTrue=False))
ws.conditional_formatting.add(f"L2:L{MAX_ROW}", FormulaRule(formula=["AND(LEN(TRIM(L2))>0,LEN(N2)=0)"], fill=red, stopIfTrue=False))

# ---------------------------------------------------------------------------
# Data tab (the only tab to publish): derived from Missionary Data
# ---------------------------------------------------------------------------
wd = wb.create_sheet("Data", index=1)
wd.append(["Organization", "Country", "State / Province", "Country Code", "Subdivision Code", "Region Code", "Website", "Notes"])
for c in wd[1]:
    c.font = Font(bold=True)
MD = "'Missionary Data'"
for r in range(2, MAX_ROW + 1):
    wd.cell(row=r, column=1, value=f'=IF(LEN({MD}!J{r}&"")>0,{MD}!J{r}&"",{MD}!B{r}&"")')
    for col, src_col in zip((2, 3, 4, 5, 6), "KLMNO"):
        wd.cell(row=r, column=col, value=f'={MD}!{src_col}{r}&""')
wd.freeze_panes = "A2"
for col, width in zip("ABCDEFGH", (34, 24, 20, 12, 15, 12, 12, 12)):
    wd.column_dimensions[col].width = width

# ---------------------------------------------------------------------------
# Instructions tab: append a section for the map columns
# ---------------------------------------------------------------------------
wi = wb["Instructions & Notes"]
last = max((c.row for row in wi.iter_rows() for c in row if c.value not in (None, "")), default=1)
title_font = wi.cell(row=3, column=1).font.copy()
lines = [
    ("MAP COLUMNS (J TO O) AND THE PUBLISHED 'Data' TAB", None),
    ("Map Label", "The name shown publicly on the website map. Defaults to the organization; remove missionaries' personal names here if they should not appear online."),
    ("Country", "Pick from the dropdown. Leave blank for global ministries; they are simply not placed on the map. Territories such as Puerto Rico or Hong Kong are their own entries."),
    ("State / Province", "Optional. Type the state, province, or region (not the city). It turns red if the map does not know that name; blank is fine, the entry then counts toward the country."),
    ("Country Code, Subdivision Code, Region Code", "Filled by formulas. Do not type in them. Region Code is what the map uses."),
    ("Red cell", "A red Country or State / Province means the map will not find it. Re-pick the country from the dropdown or fix the spelling."),
    ("Second country", "If one ministry works in two countries and both should show, add a second row for the second country."),
    ("Data tab", "Built by formulas from this tab. It is the ONLY tab published to the website and contains no respondent names, goals, or prayer requests. Do not edit it directly."),
    ("Countries / Subdivisions tabs", "Lookup lists generated from the map data. Do not edit; they are replaced when the map data is updated."),
    ("Publishing", "File > Share > Publish to web > choose the 'Data' tab only > Comma-separated values (.csv) > Publish. The map reads that link; changes appear within a few minutes."),
]
row = last + 2
for k, v in lines:
    wi.cell(row=row, column=1, value=k)
    if v is None:
        wi.cell(row=row, column=1).font = title_font
    else:
        wi.cell(row=row, column=2, value=v)
        wi.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    row += 1

os.makedirs(OUT_DIR, exist_ok=True)
wb.save(OUT)
print("saved:", OUT)
print("sheets:", wb.sheetnames)
print()
print("row | organization | map label | country | state | note")
for r, org, label, country, state, note in report:
    print(f"{r:>3} | {org} | {label} | {country} | {state}{note}")
seen = {(ws.cell(row=r, column=2).value or "").strip() for r in range(2, 60)}
print()
print("plan entries not found in sheet:", [o for o in PLAN if o not in seen])
